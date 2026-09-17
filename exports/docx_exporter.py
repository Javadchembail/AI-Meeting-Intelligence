from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from backend.schemas.mom import MinutesOfMeeting


LOCAL_TIMEZONE = ZoneInfo("Asia/Kolkata")


def format_local_datetime(value) -> str:
    """Convert a UTC-aware datetime to India Standard Time for display."""

    if value is None:
        return "Not available"

    if value.tzinfo is None:
        value = value.replace(tzinfo=ZoneInfo("UTC"))

    return value.astimezone(LOCAL_TIMEZONE).strftime(
        "%d %b %Y, %I:%M %p"
    )


class DOCXExporter:
    """
    Export Minutes of Meeting into a Microsoft Word document.
    """

    def export(
        self,
        mom: MinutesOfMeeting,
        output_path: str | Path,
    ) -> Path:
        """
        Export a MinutesOfMeeting object to a DOCX file.

        Args:
            mom: Structured Minutes of Meeting.
            output_path: Destination file path.

        Returns:
            Path: Path to the generated DOCX file.
        """

        output_file = Path(output_path)

        output_file.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        document = Document()

        # --------------------------------------------------
        # Title
        # --------------------------------------------------

        title = document.add_heading(
            "MINUTES OF MEETING",
            level=0,
        )

        title.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        # --------------------------------------------------
        # Meeting Information
        # --------------------------------------------------

        document.add_heading(
            "Meeting Information",
            level=1,
        )

        document.add_paragraph(
            f"Meeting: {mom.meeting_title}"
        )

        if mom.meeting_date:
            date_text = format_local_datetime(
                mom.meeting_date
            )
        else:
            date_text = "Not available"

        document.add_paragraph(
            f"Date: {date_text}"
        )

        if mom.duration_seconds is not None:
            minutes = (
                mom.duration_seconds // 60
            )
            seconds = (
                mom.duration_seconds % 60
            )

            duration_text = (
                f"{minutes:02d}:{seconds:02d}"
            )
        else:
            duration_text = "Not available"

        document.add_paragraph(
            f"Duration: {duration_text}"
        )

        # --------------------------------------------------
        # Executive Summary
        # --------------------------------------------------

        document.add_heading(
            "Executive Summary",
            level=1,
        )

        document.add_paragraph(
            mom.summary
        )

        # --------------------------------------------------
        # Key Points
        # --------------------------------------------------

        document.add_heading(
            "Key Points",
            level=1,
        )

        if mom.key_points:
            for point in mom.key_points:
                document.add_paragraph(
                    point,
                    style="List Bullet",
                )
        else:
            document.add_paragraph(
                "No key points were recorded."
            )

        # --------------------------------------------------
        # Decisions
        # --------------------------------------------------

        document.add_heading(
            "Decisions",
            level=1,
        )

        if mom.decisions:
            for decision in mom.decisions:
                document.add_paragraph(
                    decision,
                    style="List Number",
                )
        else:
            document.add_paragraph(
                "No decisions were recorded."
            )

        # --------------------------------------------------
        # Action Items
        # --------------------------------------------------

        document.add_heading(
            "Action Items",
            level=1,
        )

        if mom.action_items:
            table = document.add_table(
                rows=1,
                cols=4,
            )

            table.style = "Table Grid"

            headers = [
                "Task",
                "Assignee",
                "Deadline",
                "Priority",
            ]

            for cell, header in zip(
                table.rows[0].cells,
                headers,
            ):
                cell.text = header

                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True

            for item in mom.action_items:
                row = table.add_row().cells

                row[0].text = item.task

                row[1].text = (
                    item.assignee
                    or "Unassigned"
                )

                row[2].text = (
                    item.deadline
                    or "No deadline"
                )

                row[3].text = item.priority

        else:
            document.add_paragraph(
                "No action items were recorded."
            )

        # --------------------------------------------------
        # Next Steps
        # --------------------------------------------------

        document.add_heading(
            "Next Steps",
            level=1,
        )

        if mom.next_steps:
            for step in mom.next_steps:
                document.add_paragraph(
                    step,
                    style="List Bullet",
                )
        else:
            document.add_paragraph(
                "No next steps were recorded."
            )

        # --------------------------------------------------
        # Footer
        # --------------------------------------------------

        section = document.sections[0]

        footer = section.footer

        paragraph = footer.paragraphs[0]

        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER
        )

        run = paragraph.add_run(
            "Generated by AI Meeting Intelligence"
        )

        run.font.size = Pt(9)

        # --------------------------------------------------
        # Save
        # --------------------------------------------------

        document.save(
            output_file
        )

        return output_file
